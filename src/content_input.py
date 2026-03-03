"""
Content Input Module

Accepts presentation content through multiple methods and normalizes it
into a list of SlideContent objects ready for deck generation.

Supported input methods:
  - Structured outline (text)
  - Free-form / brain-dump text
  - Document upload (.docx, .pdf, .md, .txt)
  - Data / CSV file
  - URL / article
  - Conversational (chat)
"""

import os
import re
import csv
import json
import io
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Optional

import anthropic


class PresentationPurpose(str, Enum):
    GENERAL = "General"
    PITCH = "Pitch"
    REPORT = "Report"
    TRAINING = "Training"
    PROPOSAL = "Proposal"
    STRATEGY = "Strategy"
    BOARD = "Board"
    SALES = "Sales"
    EVENT = "Event"
    INVESTOR = "Investor"


class ContentDensity(str, Enum):
    LOW = "Low"
    MEDIUM = "Medium"
    HIGH = "High"


class SlideType(str, Enum):
    TITLE = "title"
    SECTION_DIVIDER = "section_divider"
    CONTENT_TEXT = "content_text"
    CONTENT_VISUAL = "content_visual"
    DATA_CHART = "data_chart"
    COMPARISON = "comparison"
    TIMELINE = "timeline"
    QUOTE = "quote"
    TEAM = "team"
    THANK_YOU = "thank_you"


@dataclass
class ChartData:
    chart_type: str = "bar"   # "bar", "line", "donut", "table"
    title: str = ""
    labels: list[str] = field(default_factory=list)
    series: list[dict] = field(default_factory=list)  # [{"name": ..., "values": [...]}]
    takeaway: str = ""


@dataclass
class SlideContent:
    slide_type: SlideType = SlideType.CONTENT_TEXT
    title: str = ""
    subtitle: str = ""
    body_bullets: list[str] = field(default_factory=list)
    body_text: str = ""
    speaker_notes: str = ""
    image_hint: str = ""       # description of suggested image
    chart_data: Optional[ChartData] = None
    layout_variant: str = "default"
    sequence_number: int = 0

    def to_dict(self) -> dict:
        return {
            "slide_type": self.slide_type.value,
            "title": self.title,
            "subtitle": self.subtitle,
            "body_bullets": self.body_bullets,
            "body_text": self.body_text,
            "speaker_notes": self.speaker_notes,
            "image_hint": self.image_hint,
            "chart_data": {
                "chart_type": self.chart_data.chart_type,
                "title": self.chart_data.title,
                "labels": self.chart_data.labels,
                "series": self.chart_data.series,
                "takeaway": self.chart_data.takeaway,
            } if self.chart_data else None,
            "layout_variant": self.layout_variant,
            "sequence_number": self.sequence_number,
        }


@dataclass
class DeckSpec:
    """Full specification of a presentation to be generated."""
    title: str = "Untitled Presentation"
    slides: list[SlideContent] = field(default_factory=list)
    purpose: PresentationPurpose = PresentationPurpose.GENERAL
    audience: str = "General"
    include_speaker_notes: bool = True
    language: str = "English"
    target_slide_count: Optional[int] = None

    def to_dict(self) -> dict:
        return {
            "title": self.title,
            "slides": [s.to_dict() for s in self.slides],
            "purpose": self.purpose.value,
            "audience": self.audience,
            "include_speaker_notes": self.include_speaker_notes,
            "language": self.language,
            "target_slide_count": self.target_slide_count,
        }


class ContentInputProcessor:
    """
    Converts raw user content into a structured DeckSpec using Claude AI.
    """

    def __init__(self, api_key: Optional[str] = None):
        self.client = anthropic.Anthropic(api_key=api_key or os.environ.get("ANTHROPIC_API_KEY"))

    # ------------------------------------------------------------------
    # Public input processors
    # ------------------------------------------------------------------

    def from_outline(self, outline_text: str, params: Optional[dict] = None) -> DeckSpec:
        """
        Process a structured outline where each section becomes a slide.

        Expected format:
          # Presentation Title
          ## Slide Title
          - bullet 1
          - bullet 2
          ## Another Slide
          ...
        """
        params = params or {}
        slides = self._parse_markdown_outline(outline_text)
        deck = self._build_deck_spec(slides, outline_text, params)
        return deck

    def from_freeform(self, text: str, params: Optional[dict] = None) -> DeckSpec:
        """Process unstructured notes/brain dump into a structured deck."""
        params = params or {}
        structured = self._ai_structure_content(text, params)
        return structured

    def from_document(self, file_path: str, params: Optional[dict] = None) -> DeckSpec:
        """
        Extract content from a document file (.docx, .pdf, .md, .txt).
        """
        params = params or {}
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext == ".docx":
            text = self._extract_docx(file_path)
        elif ext == ".pdf":
            text = self._extract_pdf(file_path)
        elif ext in (".md", ".txt"):
            text = path.read_text(encoding="utf-8")
        else:
            raise ValueError(f"Unsupported document format: {ext}")

        return self.from_freeform(text, params)

    def from_csv(self, csv_path: str, params: Optional[dict] = None) -> DeckSpec:
        """
        Convert a CSV/XLSX data file into chart slides.
        """
        params = params or {}
        path = Path(csv_path)

        if path.suffix.lower() == ".xlsx":
            try:
                import pandas as pd
                df = pd.read_excel(csv_path)
                rows = df.values.tolist()
                headers = list(df.columns)
            except ImportError:
                raise ImportError("pandas and openpyxl are required for .xlsx: pip install pandas openpyxl")
        else:
            with open(csv_path, newline="", encoding="utf-8") as f:
                reader = csv.reader(f)
                headers = next(reader, [])
                rows = list(reader)

        slides = self._ai_generate_chart_slides(headers, rows, params)
        deck_title = params.get("title", path.stem.replace("_", " ").title())
        return DeckSpec(
            title=deck_title,
            slides=slides,
            purpose=PresentationPurpose(params.get("purpose", "General")),
            audience=params.get("audience", "General"),
            include_speaker_notes=params.get("include_speaker_notes", True),
        )

    def from_url(self, url: str, params: Optional[dict] = None) -> DeckSpec:
        """Scrape a web page and convert its content into a deck."""
        params = params or {}
        try:
            import requests
            from bs4 import BeautifulSoup
        except ImportError:
            raise ImportError("requests and beautifulsoup4 required: pip install requests beautifulsoup4")

        headers = {"User-Agent": "Mozilla/5.0 (compatible; PPTGenerator/1.0)"}
        resp = requests.get(url, headers=headers, timeout=15)
        resp.raise_for_status()

        soup = BeautifulSoup(resp.text, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "aside"]):
            tag.decompose()
        text = soup.get_text(separator="\n", strip=True)[:8000]

        params.setdefault("title", soup.title.string if soup.title else url)
        return self.from_freeform(text, params)

    def from_conversational(self, messages: list[dict],
                             params: Optional[dict] = None) -> DeckSpec:
        """
        Build a deck from a chat conversation history.
        messages: [{"role": "user"|"assistant", "content": "..."}]
        """
        params = params or {}
        combined = "\n".join(
            f"{m['role'].upper()}: {m['content']}" for m in messages
        )
        return self.from_freeform(combined, params)

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _parse_markdown_outline(self, text: str) -> list[SlideContent]:
        """Parse a markdown-structured outline into SlideContent objects."""
        slides: list[SlideContent] = []
        current_slide: Optional[SlideContent] = None
        seq = 0

        for line in text.splitlines():
            stripped = line.strip()
            if stripped.startswith("# ") and not stripped.startswith("## "):
                # Presentation title — create title slide
                title_text = stripped[2:].strip()
                current_slide = SlideContent(
                    slide_type=SlideType.TITLE,
                    title=title_text,
                    sequence_number=seq,
                )
                slides.append(current_slide)
                seq += 1
            elif stripped.startswith("## "):
                if current_slide:
                    pass  # already appended
                slide_title = stripped[3:].strip()
                current_slide = SlideContent(
                    slide_type=SlideType.CONTENT_TEXT,
                    title=slide_title,
                    sequence_number=seq,
                )
                slides.append(current_slide)
                seq += 1
            elif stripped.startswith("### "):
                # Sub-section becomes a section divider
                current_slide = SlideContent(
                    slide_type=SlideType.SECTION_DIVIDER,
                    title=stripped[4:].strip(),
                    sequence_number=seq,
                )
                slides.append(current_slide)
                seq += 1
            elif stripped.startswith(("- ", "* ", "+ ")):
                if current_slide:
                    bullet = stripped[2:].strip()
                    current_slide.body_bullets.append(bullet)
            elif stripped.startswith("> "):
                # Blockquote becomes a quote slide
                current_slide = SlideContent(
                    slide_type=SlideType.QUOTE,
                    body_text=stripped[2:].strip(),
                    sequence_number=seq,
                )
                slides.append(current_slide)
                seq += 1
            elif stripped and current_slide:
                if current_slide.body_text:
                    current_slide.body_text += " " + stripped
                else:
                    current_slide.body_text = stripped

        return slides

    def _ai_structure_content(self, text: str, params: dict) -> DeckSpec:
        """Use Claude to parse free-form text into a structured deck."""
        purpose = params.get("purpose", "General")
        audience = params.get("audience", "General")
        target_count = params.get("target_slide_count", "Auto")
        density = params.get("content_density", "Medium")
        include_notes = params.get("include_speaker_notes", True)
        language = params.get("language", "English")
        deck_title = params.get("title", "")

        prompt = f"""You are a presentation design expert. Convert the following content into a structured slide deck.

Parameters:
- Purpose: {purpose}
- Audience: {audience}
- Target slide count: {target_count} (Auto = you decide the optimal count)
- Content density: {density} (Low = less text per slide, High = more)
- Language: {language}
- Deck title: {deck_title or "derive from content"}

Return ONLY a valid JSON object in this exact structure:
{{
  "title": "Presentation Title",
  "slides": [
    {{
      "slide_type": "title|section_divider|content_text|content_visual|data_chart|comparison|timeline|quote|team|thank_you",
      "title": "Slide Title",
      "subtitle": "Optional subtitle",
      "body_bullets": ["bullet 1", "bullet 2"],
      "body_text": "Optional paragraph text instead of bullets",
      "speaker_notes": "Speaker notes for this slide",
      "image_hint": "Description of suggested image if visual slide",
      "layout_variant": "default"
    }}
  ]
}}

Guidelines:
- Start with a title slide and end with a thank you / CTA slide
- Add section dividers between major topics
- Keep bullets short (under 10 words each) for Medium density
- Generate speaker notes that expand on the slide content
- Use data_chart type when content contains numbers/metrics
- Use comparison type for before/after or option comparisons
- Use quote type for testimonials or key statistics

Content to process:
{text[:10000]}"""

        message = self.client.messages.create(
            model="claude-opus-4-6",
            max_tokens=4096,
            messages=[{"role": "user", "content": prompt}],
        )
        raw = message.content[0].text.strip()

        # Strip markdown fences
        if raw.startswith("```"):
            raw = re.sub(r"^```[a-z]*\n?", "", raw)
            raw = re.sub(r"\n?```$", "", raw)

        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            # Fallback: create a simple deck with the raw text
            return DeckSpec(
                title=deck_title or "Presentation",
                slides=[
                    SlideContent(slide_type=SlideType.TITLE, title=deck_title or "Presentation"),
                    SlideContent(slide_type=SlideType.CONTENT_TEXT,
                                 title="Content", body_text=text[:500]),
                ],
            )

        slides = []
        for i, s in enumerate(data.get("slides", [])):
            try:
                slide_type = SlideType(s.get("slide_type", "content_text"))
            except ValueError:
                slide_type = SlideType.CONTENT_TEXT

            chart_data = None
            if s.get("chart_data"):
                cd = s["chart_data"]
                chart_data = ChartData(
                    chart_type=cd.get("chart_type", "bar"),
                    title=cd.get("title", ""),
                    labels=cd.get("labels", []),
                    series=cd.get("series", []),
                    takeaway=cd.get("takeaway", ""),
                )

            slides.append(SlideContent(
                slide_type=slide_type,
                title=s.get("title", ""),
                subtitle=s.get("subtitle", ""),
                body_bullets=s.get("body_bullets", []),
                body_text=s.get("body_text", ""),
                speaker_notes=s.get("speaker_notes", "") if include_notes else "",
                image_hint=s.get("image_hint", ""),
                chart_data=chart_data,
                layout_variant=s.get("layout_variant", "default"),
                sequence_number=i,
            ))

        return DeckSpec(
            title=data.get("title", deck_title or "Presentation"),
            slides=slides,
            purpose=PresentationPurpose(purpose) if purpose in PresentationPurpose._value2member_map_ else PresentationPurpose.GENERAL,
            audience=audience,
            include_speaker_notes=include_notes,
            language=language,
            target_slide_count=target_count if isinstance(target_count, int) else None,
        )

    def _build_deck_spec(self, slides: list[SlideContent], original_text: str,
                         params: dict) -> DeckSpec:
        """Enhance parsed outline slides with speaker notes via Claude."""
        include_notes = params.get("include_speaker_notes", True)
        if not include_notes:
            return DeckSpec(
                title=slides[0].title if slides and slides[0].slide_type == SlideType.TITLE else "Presentation",
                slides=slides,
                purpose=PresentationPurpose(params.get("purpose", "General")),
                audience=params.get("audience", "General"),
            )

        # Generate speaker notes for each content slide
        for slide in slides:
            if slide.slide_type in (SlideType.TITLE, SlideType.THANK_YOU,
                                    SlideType.SECTION_DIVIDER):
                continue
            bullets_text = "\n".join(f"- {b}" for b in slide.body_bullets)
            note_prompt = (
                f"Write concise speaker notes (2-3 sentences) for a presentation slide.\n"
                f"Slide title: {slide.title}\n"
                f"Bullets:\n{bullets_text}\n"
                f"Notes should expand on the bullets with context and talking points."
            )
            try:
                msg = self.client.messages.create(
                    model="claude-haiku-4-5-20251001",
                    max_tokens=256,
                    messages=[{"role": "user", "content": note_prompt}],
                )
                slide.speaker_notes = msg.content[0].text.strip()
            except Exception:
                pass

        deck_title = slides[0].title if slides and slides[0].slide_type == SlideType.TITLE else params.get("title", "Presentation")
        return DeckSpec(
            title=deck_title,
            slides=slides,
            purpose=PresentationPurpose(params.get("purpose", "General")),
            audience=params.get("audience", "General"),
            include_speaker_notes=include_notes,
        )

    def _ai_generate_chart_slides(self, headers: list[str], rows: list[list],
                                   params: dict) -> list[SlideContent]:
        """Generate chart-focused slides from tabular data."""
        # Summarize the data for Claude
        preview_rows = rows[:10]
        data_preview = "Headers: " + ", ".join(str(h) for h in headers) + "\n"
        for row in preview_rows:
            data_preview += " | ".join(str(v) for v in row) + "\n"
        data_preview += f"(Total rows: {len(rows)})"

        prompt = (
            f"You are a data visualization expert. Given this tabular data, "
            f"create slide specifications for chart slides.\n\n"
            f"Data:\n{data_preview}\n\n"
            f"Return JSON array of slide objects with this structure:\n"
            f'[{{"slide_type": "data_chart", "title": "...", "chart_data": '
            f'{{"chart_type": "bar|line|donut", "title": "...", "labels": [...], '
            f'"series": [{{"name": "...", "values": [...]}}], "takeaway": "key insight"}}, '
            f'"speaker_notes": "..."}}]\n\n'
            f"Create 1-3 chart slides that best represent the data insights."
        )

        try:
            msg = self.client.messages.create(
                model="claude-opus-4-6",
                max_tokens=2048,
                messages=[{"role": "user", "content": prompt}],
            )
            raw = msg.content[0].text.strip()
            if raw.startswith("```"):
                raw = re.sub(r"^```[a-z]*\n?", "", raw)
                raw = re.sub(r"\n?```$", "", raw)
            slides_data = json.loads(raw)
        except Exception:
            # Fallback: create a simple data slide
            return [SlideContent(
                slide_type=SlideType.DATA_CHART,
                title="Data Overview",
                body_text=f"Data with columns: {', '.join(str(h) for h in headers)}",
            )]

        slides = []
        for i, s in enumerate(slides_data):
            cd = s.get("chart_data", {})
            slides.append(SlideContent(
                slide_type=SlideType.DATA_CHART,
                title=s.get("title", "Data"),
                speaker_notes=s.get("speaker_notes", ""),
                chart_data=ChartData(
                    chart_type=cd.get("chart_type", "bar"),
                    title=cd.get("title", ""),
                    labels=cd.get("labels", []),
                    series=cd.get("series", []),
                    takeaway=cd.get("takeaway", ""),
                ),
                sequence_number=i,
            ))
        return slides

    def _extract_docx(self, path: str) -> str:
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx is required: pip install python-docx")
        doc = Document(path)
        return "\n".join(para.text for para in doc.paragraphs if para.text.strip())

    def _extract_pdf(self, path: str) -> str:
        try:
            import pdfplumber
        except ImportError:
            raise ImportError("pdfplumber is required: pip install pdfplumber")
        texts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages[:50]:
                t = page.extract_text()
                if t:
                    texts.append(t)
        return "\n\n".join(texts)
